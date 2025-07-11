import streamlit as st
import pandas as pd
import matplotlib.pyplot as plt
import base64
from io import BytesIO
from fpdf import FPDF
import os

# === Affichage du logo ===
chemin_logo = os.path.join(os.path.dirname(__file__), "interface.jpg")
if os.path.exists(chemin_logo):
    st.image(chemin_logo, width=300)

st.title("📊 Notabilis - Analyse des notes scolaires")

uploaded_file = st.file_uploader("📁 Déposez un fichier (.csv, .txt, .docx)", type=["csv", "txt", "docx"])

if uploaded_file:
    try:
        if uploaded_file.name.endswith(".csv") or uploaded_file.name.endswith(".txt"):
            df = pd.read_csv(uploaded_file)
        elif uploaded_file.name.endswith(".docx"):
            from docx import Document
            doc = Document(uploaded_file)
            data = [line.text.split("\t") for line in doc.paragraphs if line.text.strip()]
            df = pd.DataFrame(data[1:], columns=data[0])
        else:
            st.error("❌ Format non pris en charge.")
            st.stop()

        df.columns = [col.strip().lower() for col in df.columns]

        # Recherche colonnes
        nom_col = next((col for col in df.columns if "nom" in col), None)
        devoir_cols = [col for col in df.columns if "devoir" in col]
        interro_cols = [col for col in df.columns if "interro" in col]
        coef_col = next((col for col in df.columns if "coef" in col), None)

        if not (nom_col and devoir_cols and interro_cols and coef_col):
            st.error("❌ Certaines colonnes sont manquantes : 'Nom', 'Interro', 'Devoir', ou 'Coef'")
            st.stop()

        df[nom_col] = df[nom_col].astype(str)
        df[interro_cols + devoir_cols + [coef_col]] = df[interro_cols + devoir_cols + [coef_col]].apply(pd.to_numeric, errors="coerce").fillna(0)

        # Calcul des moyennes
        df["moy_interro"] = df[interro_cols].mean(axis=1)
        df["moyenne"] = (df["moy_interro"] + df[devoir_cols[0]] + df[devoir_cols[1]]) / 3
        df["moy_coeff"] = df["moyenne"] * df[coef_col]

        # Rang
        df["rang"] = df["moyenne"].rank(ascending=False, method='min').astype(int)

        # Statistiques globales
        moyenne_classe = round(df["moyenne"].mean(), 2)
        max_moyenne = round(df["moyenne"].max(), 2)
        min_moyenne = round(df["moyenne"].min(), 2)

        st.success(f"🎯 Moyenne de la classe : {moyenne_classe} | 🏅 Max : {max_moyenne} | 🔻 Min : {min_moyenne}")
        st.dataframe(df[[nom_col, *interro_cols, *devoir_cols, coef_col, "moy_interro", "moyenne", "moy_coeff", "rang"]])

        # === Diagramme des moyennes ===
        fig, ax = plt.subplots()
        ax.bar(df[nom_col], df["moyenne"], color='skyblue')
        plt.xticks(rotation=90)
        plt.title("Diagramme des moyennes")
        st.pyplot(fig)

        # === Génération du PDF ===
        def generate_pdf(dataframe):
            pdf = FPDF()
            pdf.add_page()
            pdf.set_font("Arial", size=12)
            pdf.cell(200, 10, txt="Rapport de notes - Notabilis", ln=True, align='C')
            pdf.ln(10)

            for i, row in dataframe.iterrows():
                ligne = f"{row[nom_col]} | Moyenne: {round(row['moyenne'], 2)} | Rang: {row['rang']}"
                pdf.cell(200, 10, txt=ligne, ln=True)

            pdf.ln(10)
            pdf.cell(200, 10, txt=f"Moyenne de la classe: {moyenne_classe}", ln=True)
            pdf.cell(200, 10, txt=f"Plus forte moyenne: {max_moyenne}", ln=True)
            pdf.cell(200, 10, txt=f"Plus faible moyenne: {min_moyenne}", ln=True)

            return pdf.output(dest='S').encode('latin1')  # ✅ Correction du bug ici

        # === Bouton de téléchargement du rapport PDF ===
        pdf_bytes = generate_pdf(df)
        b64 = base64.b64encode(pdf_bytes).decode()
        href = f'<a href="data:application/octet-stream;base64,{b64}" download="rapport_notabilis.pdf">📄 Télécharger le rapport PDF</a>'
        st.markdown(href, unsafe_allow_html=True)

    except Exception as e:
        st.error(f"❌ Une erreur est survenue : {e}")